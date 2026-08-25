"""
sobre_a_plantillas.py — Fase 4c: llenado sobre las PLANTILLAS OFICIALES.

Objetivo (pedido del cliente): el formulario generado debe ser IDÉNTICO al
oficial en formato y contenido — escudo nacional, recuadros, tablas, tipografía
— con solo estas transformaciones en tiempo de ejecución:

  1. LOGO institucional FUERA: el logo vive en un content control con alias
     "Logo de la dependencia gubernamental" → se vacía ese SDT quirúrgicamente.
     El escudo nacional es una imagen normal del cuerpo y NO se toca.
  2. INSTITUCIÓN parametrizada: el texto "INSTITUTO NACIONAL DE AGUAS POTABLES
     Y ALCANTARILLADOS"/"INAPA" se sustituye por la institución del proceso.
  3. TEXTOS ROJOS: son instrucciones "[indicar…]"/"(poner aquí…)". Se
     reemplazan por el dato real EN NEGRO; los que son pura instrucción sin
     dato (p. ej. el bloque "[El Oferente deberá completar…]") se eliminan.
     Si el dato no existe en el expediente, queda el hueco sombreado
     [[PENDIENTE: campo]] (regla de Fase 4).
  4. SDTs: "No. del Expediente" → referencia; "Fecha de emisión" → hoy.
  5. TABLAS de datos (F.036 equipos, D.049 experiencia): se llenan clonando
     el formato de la fila de la plantilla.
  6. FIRMA: si el firmante principal tiene firma_url, la imagen se inserta
     sobre la línea "Firma ____".

Las plantillas viven en plantillas_sncc/{clave}.docx (binarios versionados en
el repo). Si falta la plantilla de una clave, sobre_a_core cae a su generador
programático (fallback, nunca rompe).
"""

from __future__ import annotations

import copy
import io
import os
import re
from datetime import date

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

DIR_PLANTILLAS = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                              "plantillas_sncc")

MARCA = "[[PENDIENTE: {campo}]]"
ROJO = "FF0000"
# Nombre/siglas de la institución de ORIGEN de las plantillas (se sustituyen)
_INST_ORIGEN = ("INSTITUTO NACIONAL DE AGUAS POTABLES Y ALCANTARILLADOS",
                "INAPA")

W = qn("w:t")


def hay_plantilla(clave: str) -> bool:
    return os.path.exists(os.path.join(DIR_PLANTILLAS, f"{clave}.docx"))


def _abrir(clave: str) -> Document:
    return Document(os.path.join(DIR_PLANTILLAS, f"{clave}.docx"))


def _v(reg: dict | None, *campos):
    if not reg:
        return None
    for c in campos:
        val = reg.get(c)
        if val is not None and str(val).strip():
            return str(val).strip()
    return None


def _dato(reg, etiqueta, *campos):
    return _v(reg, *campos) or MARCA.format(campo=etiqueta)


# ══════════════════════════════════════════════════════════════
# Transformaciones genéricas
# ══════════════════════════════════════════════════════════════

def _quitar_logo_institucion(doc: Document) -> None:
    """
    Vacía SOLO la imagen de los SDT 'Logo de la dependencia gubernamental'.
    OJO: dentro de ese SDT también viven text boxes (w:drawing con
    w:txbxContent) que traen el nombre de la institución y el título del
    formulario — esos se conservan; solo se elimina el drawing que es
    imagen pura (tiene a:blip y no tiene txbxContent). El escudo nacional
    está fuera del SDT y no se toca.
    """
    A_BLIP = qn("a:blip")
    TXBX = qn("w:txbxContent")
    for sdt in doc.element.body.iter(qn("w:sdt")):
        props = sdt.find(qn("w:sdtPr"))
        alias = props.find(qn("w:alias")) if props is not None else None
        val = alias.get(qn("w:val")) if alias is not None else ""
        if val and "logo" in val.lower():
            contenido = sdt.find(qn("w:sdtContent"))
            if contenido is None:
                continue
            for dib in list(contenido.iter(qn("w:drawing"))):
                es_imagen = dib.find(f".//{A_BLIP}") is not None
                tiene_texto = dib.find(f".//{TXBX}") is not None
                if es_imagen and not tiene_texto:
                    dib.getparent().remove(dib)
            for pic in list(contenido.iter(qn("w:pict"))):
                if pic.find(f".//{TXBX}") is None:
                    pic.getparent().remove(pic)


def _set_texto_run(run, texto: str) -> None:
    """
    Escribe texto en un run tocando SOLO sus w:t. Jamás usar run.text=…:
    ese setter borra TODOS los hijos del run, incluidos los w:drawing de
    text boxes anclados (así se perdían el título y la institución).
    """
    ts = run._element.findall(qn("w:t"))
    if ts:
        ts[0].text = texto
        for t in ts[1:]:
            t.text = ""
    elif texto:
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = texto
        run._element.append(t)


def _llenar_placeholders(doc: Document, referencia: str) -> None:
    """
    No. de expediente y fecha: se llenan por el TEXTO del placeholder
    ("Click here to enter text." / "Seleccione la fecha"), que es infalible
    aunque el valor viva dentro de text boxes anidados en los SDT.
    """
    hoy = date.today().strftime("%d/%m/%Y")
    for t in doc.element.body.iter(W):
        txt = (t.text or "").lower()
        if "click here to enter" in txt or "haga clic aqu" in txt:
            t.text = referencia
        elif "seleccione la fecha" in txt:
            t.text = hoy


# compat: el pipeline llama _llenar_sdts
_llenar_sdts = _llenar_placeholders


def _sustituir_institucion(doc: Document, institucion: str | None) -> None:
    """
    Cambia el nombre/siglas de la institución de origen por la del proceso.
    Trabaja a nivel XML sobre CADA párrafo (w:p) del documento y sus headers,
    lo que cubre también los text boxes (w:txbxContent) donde las plantillas
    oficiales ponen el nombre — y aguanta texto partido en varios runs.
    """
    inst = (institucion or MARCA.format(campo="Institución del proceso")).upper()

    def _procesar_raiz(raiz):
        TXBX = qn("w:txbxContent")
        for p in raiz.iter(qn("w:p")):
            # SOLO párrafos hoja: los que anclan un text box contienen
            # anidados los párrafos del text box (y su copia de
            # compatibilidad mc:Fallback) — procesarlos duplicaría el texto.
            if p.find(f".//{TXBX}") is not None:
                continue
            ts = [t for t in p.iter(W)]
            if not ts:
                continue
            texto = "".join(t.text or "" for t in ts)
            mayus = texto.upper()
            tiene_nombre = _INST_ORIGEN[0] in mayus
            tiene_siglas = bool(re.search(rf"\b{_INST_ORIGEN[1]}\b", mayus))
            if not tiene_nombre and not tiene_siglas:
                continue
            if tiene_nombre:
                nuevo = re.sub(re.escape(_INST_ORIGEN[0]), inst, texto,
                               flags=re.IGNORECASE)
                # las siglas que acompañen al nombre sobran → fuera
                nuevo = re.sub(rf"\s*\b{_INST_ORIGEN[1]}\b", "", nuevo,
                               flags=re.IGNORECASE)
            elif texto.strip().upper() == _INST_ORIGEN[1]:
                nuevo = ""                     # línea de siglas sola → fuera
            else:
                nuevo = re.sub(rf"\b{_INST_ORIGEN[1]}\b", inst, texto,
                               flags=re.IGNORECASE)
            ts[0].text = nuevo
            for t in ts[1:]:
                t.text = ""

    _procesar_raiz(doc.element.body)
    for sec in doc.sections:
        _procesar_raiz(sec.header._element)
        _procesar_raiz(sec.footer._element)


def _reescribir(parrafo, texto: str) -> None:
    """Reemplaza el texto de un párrafo conservando el formato del 1er run."""
    if not parrafo.runs:
        parrafo.add_run(texto)
        return
    _set_texto_run(parrafo.runs[0], texto)
    for r in parrafo.runs[1:]:
        _set_texto_run(r, "")


def _es_rojo(run) -> bool:
    c = run.font.color
    return bool(c and c.rgb and str(c.rgb) == ROJO)


def _pintar_valor(run, valor: str) -> None:
    """Run rojo → dato real en negro, sin cursiva; hueco → sombreado."""
    _set_texto_run(run, valor)
    run.font.color.rgb = None            # hereda (negro/auto)
    run.font.italic = False
    run.font.bold = False
    if "[[PENDIENTE" in valor:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "FFF2CC")
        run._element.get_or_add_rPr().append(shd)


def _llenar_rojos(doc: Document, mapa: list[tuple[str, str]]) -> None:
    """
    Recorre body + tablas. Agrupa runs rojos consecutivos de cada párrafo,
    junta su texto y:
      - si alguna palabra clave del mapa aparece → 1er run = valor en negro,
        el resto se vacía;
      - si no matchea nada → es instrucción pura: se elimina el texto rojo.
    El mapa es [(palabra_clave_en_minusculas, valor)], primera coincidencia gana.
    """
    def _proc(parrafos):
        for p in parrafos:
            grupo = []
            for r in list(p.runs) + [None]:          # None = cierre del último grupo
                if r is not None and _es_rojo(r):
                    grupo.append(r)
                    continue
                if grupo:
                    texto = "".join(g.text or "" for g in grupo).lower()
                    valor = next((v for k, v in mapa if k in texto), None)
                    if valor is not None:
                        _pintar_valor(grupo[0], valor)
                        for g in grupo[1:]:
                            _set_texto_run(g, "")
                    else:                             # instrucción pura → fuera
                        for g in grupo:
                            _set_texto_run(g, "")
                    grupo = []

    _proc(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                _proc(cell.paragraphs)


def _quitar_frase(doc: Document, frase: str) -> None:
    """Elimina una frase literal residual (texto negro de instrucción que
    acompaña a un rojo ya reemplazado, p. ej. 'la Entidad')."""
    TXBX = qn("w:txbxContent")
    for p in doc.element.body.iter(qn("w:p")):
        if p.find(f".//{TXBX}") is not None:
            continue
        ts = [t for t in p.iter(W)]
        texto = "".join(t.text or "" for t in ts)
        if frase in texto:
            nuevo = texto.replace(frase, "")
            ts[0].text = nuevo
            for t in ts[1:]:
                t.text = ""


def _insertar_firma(doc: Document, firma_png: bytes | None) -> None:
    if not firma_png:
        return
    for p in doc.paragraphs:
        if p.text.strip().startswith("Firma"):
            try:
                nuevo = p.insert_paragraph_before()
                nuevo.add_run().add_picture(io.BytesIO(firma_png),
                                            width=Inches(1.8))
            except Exception:
                pass
            return


def _llenar_nombre_calidad(doc: Document, nombre: str, cargo: str,
                           razon: str) -> None:
    """Línea '(Nombre y apellido) ____ en calidad de ____ … representación de ____'."""
    for p in doc.paragraphs:
        if "Nombre y apellido" in p.text and "_" in p.text:
            texto = p.text
            partes = re.split(r"_{4,}", texto)
            valores = [nombre, cargo, razon]
            nuevo, i = partes[0], 0
            for parte in partes[1:]:
                nuevo += (valores[i] if i < len(valores) else "____") + parte
                i += 1
            _reescribir(p, nuevo)
            return


def _clonar_fila(tabla, indice_modelo: int) -> object:
    """Duplica una fila conservando su formato; devuelve la fila nueva."""
    modelo = tabla.rows[indice_modelo]._tr
    nueva = copy.deepcopy(modelo)
    modelo.addnext(nueva)
    return tabla.rows[indice_modelo + 1]


def _set_celda(celda, valor) -> None:
    txt = "" if valor is None else str(valor)
    if celda.paragraphs and celda.paragraphs[0].runs:
        _reescribir(celda.paragraphs[0], txt)
        for p in celda.paragraphs[1:]:
            _reescribir(p, "")
    else:
        celda.text = txt


def _llenar_tabla(tabla, filas: list[list], fila_datos: int = 1) -> None:
    """
    Llena una tabla de plantilla: usa las filas vacías existentes a partir de
    fila_datos y clona más (con el mismo formato) si hacen falta.
    """
    disponibles = len(tabla.rows) - fila_datos
    while disponibles < len(filas):
        _clonar_fila(tabla, fila_datos)
        disponibles += 1
    for i, datos in enumerate(filas):
        celdas = tabla.rows[fila_datos + i].cells
        for j, val in enumerate(datos[:len(celdas)]):
            _set_celda(celdas[j], val)


def _bytes(doc: Document) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _base(clave: str, ctx: dict, mapa_rojos: list[tuple[str, str]]) -> Document:
    """Pipeline común: abrir plantilla → logo fuera → institución → SDT → rojos."""
    doc = _abrir(clave)
    _quitar_logo_institucion(doc)
    _sustituir_institucion(doc, ctx.get("institucion"))
    _llenar_sdts(doc, ctx["referencia"])
    _llenar_rojos(doc, mapa_rojos)
    return doc


def _mapa_comun(ctx: dict) -> list[tuple[str, str]]:
    emp, fir, proc = ctx["empresa"], ctx["firmante"], ctx["proceso"]
    razon = _dato(emp, "Razón social", "razon_social", "nombre_perfil")
    objeto = _v(proc, "nombre_proceso") or MARCA.format(campo="Objeto del proceso")
    return [
        # instrucciones con dato — primera coincidencia gana
        ("nombre jurídico de cada miembro", "N/A"),
        ("nombre jurídico del oferente", razon),
        ("poner aquí nombre del oferente", razon),
        ("nombre del oferente", razon),
        # "Indicar Nombre de <la Entidad>": es la institución contratante
        ("indicar nombre de", (ctx.get("institucion")
                               or MARCA.format(campo="Institución")).upper() + " "),
        ("registro de proveedores del estado", _dato(emp, "RPE", "rpe")),
        ("nombre del representante autorizado",
         _dato(fir, "Nombre del representante", "nombre_completo")),
        ("dirección del representante", _dato(fir, "Dirección", "direccion")),
        ("teléfono y fax del representante", _dato(fir, "Teléfono", "telefono")),
        ("correo electrónico del representante", _dato(fir, "Correo", "email")),
        ("denominación de la obra", f"{objeto} — Ref. {ctx['referencia']} "),
        ("procedimiento de contratación", f"{ctx['referencia']} "),
        ("incluir en números", ""),
        # instrucciones puras (sin dato) que deben desaparecer las cubre
        # el comportamiento por defecto de _llenar_rojos.
    ]


# ══════════════════════════════════════════════════════════════
# Generadores por plantilla oficial
# ══════════════════════════════════════════════════════════════

def plantilla_f034(sb, eid: str, ctx: dict) -> bytes:
    doc = _base("f034", ctx, _mapa_comun(ctx))
    fir, emp = ctx["firmante"], ctx["empresa"]
    _llenar_nombre_calidad(
        doc,
        _dato(fir, "Nombre del representante", "nombre_completo"),
        _v(fir, "cargo") or "Representante Legal",
        _dato(emp, "Razón social", "razon_social", "nombre_perfil"))
    _insertar_firma(doc, ctx.get("firma_png"))
    return _bytes(doc)


def plantilla_f042(sb, eid: str, ctx: dict) -> bytes:
    emp, fir = ctx["empresa"], ctx["firmante"]
    mapa = _mapa_comun(ctx) + [
        ("rnc", _dato(emp, "RNC", "rnc")),
        ("cédula/ pasaporte", _dato(emp, "RNC", "rnc")),
        ("domicilio legal", _dato(emp, "Dirección", "direccion_completa",
                                  "domicilio", "direccion")),
    ]
    doc = _base("f042", ctx, mapa)
    # celdas del F.042 que van tras los dos puntos y no tienen texto rojo
    _rellenar_tras_etiqueta(doc, "RNC/ Cédula/ Pasaporte", _dato(emp, "RNC", "rnc"))
    _rellenar_tras_etiqueta(doc, "Domicilio legal",
                            _dato(emp, "Dirección", "direccion_completa",
                                  "domicilio", "direccion"))
    _insertar_firma(doc, ctx.get("firma_png"))
    return _bytes(doc)


def _rellenar_tras_etiqueta(doc: Document, etiqueta: str, valor: str) -> None:
    """Celdas tipo '3. RNC/…:' sin placeholder rojo: agrega el valor al final."""
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                txt = cell.text
                if etiqueta.lower() in txt.lower() and valor not in txt:
                    partes = txt.rstrip().rstrip(":")
                    _set_celda(cell, f"{partes}: {valor}")
                    return


def plantilla_f035(sb, eid: str, ctx: dict) -> bytes:
    doc = _base("f035", ctx, _mapa_comun(ctx))
    _insertar_firma(doc, ctx.get("firma_png"))
    return _bytes(doc)


def plantilla_f036(sb, eid: str, ctx: dict) -> bytes:
    doc = _base("f036", ctx, _mapa_comun(ctx))
    _quitar_frase(doc, "la Entidad")
    equipos = (sb.table("bid_equipos").select("*")
               .eq("empresa_id", eid).eq("activo", True)
               .order("descripcion").execute().data or [])
    if doc.tables and equipos:
        antiguedad = lambda e: (date.today().year - int(e["anio"])) \
            if str(e.get("anio") or "").isdigit() else ""
        filas = [[i + 1,
                  " ".join(x for x in [_v(e, "marca"), _v(e, "descripcion")] if x),
                  _v(e, "capacidad") or "",
                  e.get("cantidad") or 1,
                  antiguedad(e),
                  "P" if (_v(e, "propiedad") or "").lower().startswith("prop") else "A",
                  "",
                  ""]
                 for i, e in enumerate(equipos)]
        _llenar_tabla(doc.tables[0], filas, fila_datos=1)
    _insertar_firma(doc, ctx.get("firma_png"))
    return _bytes(doc)


def plantilla_f037(sb, eid: str, ctx: dict) -> bytes:
    doc = _base("f037", ctx, _mapa_comun(ctx))
    _quitar_frase(doc, "la Entidad")
    personal = (sb.table("bid_personal").select("*")
                .eq("empresa_id", eid).eq("activo", True)
                .order("nombre_completo").execute().data or [])
    if doc.tables and personal:
        filas = [[i + 1,
                  _v(p, "nombre_completo") or "",
                  _v(p, "cargo_empresa") or "",
                  _v(p, "profesion") or "",
                  p.get("experiencia_general_anios") or ""]
                 for i, p in enumerate(personal)]
        _llenar_tabla(doc.tables[0], filas, fila_datos=1)
    elif personal:
        # la plantilla no trae tabla: se agrega una antes del bloque de firma
        _tabla_apendice(doc, ["No.", "Nombre completo", "Cargo", "Profesión",
                              "Años exp."],
                        [[i + 1, _v(p, "nombre_completo") or "",
                          _v(p, "cargo_empresa") or "", _v(p, "profesion") or "",
                          p.get("experiencia_general_anios") or ""]
                         for i, p in enumerate(personal)])
    _insertar_firma(doc, ctx.get("firma_png"))
    return _bytes(doc)


def _tabla_apendice(doc: Document, encabezados: list[str], filas: list[list]):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    for i, h in enumerate(encabezados):
        t.rows[0].cells[i].text = h
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            r.bold = True
    for f in filas:
        celdas = t.add_row().cells
        for i, v in enumerate(f[:len(celdas)]):
            celdas[i].text = "" if v is None else str(v)
    # mover la tabla justo antes de la línea de firma si existe
    for p in doc.paragraphs:
        if p.text.strip().startswith("Firma"):
            p._element.addprevious(t._element)
            return


def plantilla_d044(sb, eid: str, ctx: dict) -> bytes:
    doc = _base("d044", ctx, _mapa_comun(ctx))
    _insertar_firma(doc, ctx.get("firma_png"))
    return _bytes(doc)


def plantilla_d049(sb, eid: str, ctx: dict) -> bytes:
    doc = _base("d049", ctx, _mapa_comun(ctx))
    exps = (sb.table("bid_experiencia").select("*")
            .eq("empresa_id", eid)
            .order("fecha_inicio", desc=True).execute().data or [])
    if doc.tables and exps:
        filas = []
        for e in exps:
            ini, fin = str(e.get("fecha_inicio") or "")[:10], str(e.get("fecha_fin") or "")[:10]
            filas.append([
                _v(e, "nombre_proyecto") or "",
                f"{_v(e, 'moneda') or 'RD$'} {e.get('monto_contrato') or ''}".strip(),
                f"{ini} — {fin}".strip(" —"),
                ini,
                "100%" if (_v(e, "estado") or "").lower().startswith("termin") else
                (_v(e, "estado") or ""),
                _v(e, "cliente") or "",
                _v(e, "ubicacion", "provincia") or "",
                _v(e, "tipo_obra") or ""])
        _llenar_tabla(doc.tables[0], filas, fila_datos=1)
    _insertar_firma(doc, ctx.get("firma_png"))
    return _bytes(doc)


def _por_persona(clave: str, sb, eid: str, ctx: dict,
                 mapa_persona) -> list[tuple[str, bytes]]:
    """Un docx por persona (D.045/D.048): plantilla oficial llenada por técnico."""
    personal = (sb.table("bid_personal").select("*")
                .eq("empresa_id", eid).eq("activo", True)
                .order("nombre_completo").execute().data or [])
    salida = []
    for p in personal or [{}]:
        doc = _base(clave, ctx, _mapa_comun(ctx) + mapa_persona(p))
        firma_p = None
        if p.get("firma_url"):
            try:
                from sobre_a_core import _descargar_binario
                firma_p = _descargar_binario(sb, p["firma_url"])
            except Exception:
                firma_p = None
        _insertar_firma(doc, firma_p or ctx.get("firma_png"))
        sufijo = re.sub(r"[^A-Za-z0-9]+", "_", _v(p, "nombre_completo") or "personal")[:40]
        salida.append((sufijo, _bytes(doc)))
    return salida


def plantilla_d045(sb, eid: str, ctx: dict) -> list[tuple[str, bytes]]:
    def mapa(p):
        return [
            ("nombre del personal", _dato(p, "Nombre", "nombre_completo")),
            ("nombre completo", _dato(p, "Nombre", "nombre_completo")),
            ("cargo", _dato(p, "Cargo", "cargo_empresa")),
            ("profesión", _dato(p, "Profesión", "profesion")),
            ("fecha de nacimiento", _v(p, "fecha_nacimiento") or ""),
            ("cédula", _dato(p, "Cédula", "cedula")),
        ]
    return _por_persona("d045", sb, eid, ctx, mapa)


def plantilla_d048(sb, eid: str, ctx: dict) -> list[tuple[str, bytes]]:
    def mapa(p):
        return [
            ("nombre del personal", _dato(p, "Nombre", "nombre_completo")),
            ("nombre completo", _dato(p, "Nombre", "nombre_completo")),
            ("cargo", _dato(p, "Cargo", "cargo_empresa")),
            ("profesión", _dato(p, "Profesión", "profesion")),
        ]
    return _por_persona("d048", sb, eid, ctx, mapa)


def plantilla_etic(sb, eid: str, ctx: dict) -> bytes:
    emp, fir = ctx["empresa"], ctx["firmante"]
    mapa = _mapa_comun(ctx) + [
        ("cédula de identidad", _dato(fir, "Cédula", "cedula")),
        ("nacionalidad", _v(fir, "nacionalidad") or "dominicana"),
    ]
    doc = _base("etic", ctx, mapa)
    _insertar_firma(doc, ctx.get("firma_png"))
    return _bytes(doc)


# Registro: clave → generador sobre plantilla oficial
PLANTILLA_GENERADORES = {
    "f034": plantilla_f034, "f042": plantilla_f042, "f035": plantilla_f035,
    "f036": plantilla_f036, "f037": plantilla_f037, "d044": plantilla_d044,
    "d045": plantilla_d045, "d048": plantilla_d048, "d049": plantilla_d049,
    "etic": plantilla_etic,
}
